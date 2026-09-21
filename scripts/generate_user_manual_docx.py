from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "QuickOps快维用户使用手册.md"
OUTPUT = ROOT / "docs" / "QuickOps快维用户使用手册.docx"
TEAL = "0C8F88"
DARK = "173039"
MUTED = "607780"
PALE = "E6F3F2"
BODY_FONT = "Hiragino Sans GB"
MONO_CJK_FONT = "Hiragino Sans GB"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = BODY_FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, 9.3)
            run.font.name = "SFMono-Regular"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_CJK_FONT)
            run.font.color.rgb = RGBColor.from_string("275C62")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 8.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, 8.5)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.3

    for name, size, color, before, after in (
        ("Title", 28, DARK, 0, 8),
        ("Subtitle", 11, MUTED, 0, 4),
        ("Heading 1", 17, DARK, 16, 7),
        ("Heading 2", 12.5, TEAL, 11, 4),
        ("Heading 3", 10.8, DARK, 8, 3),
    ):
        style = document.styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_cover(document: Document) -> None:
    for _ in range(4):
        document.add_paragraph()
    line = document.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    marker = line.add_run("●")
    marker.font.color.rgb = RGBColor.from_string(TEAL)
    marker.font.size = Pt(14)
    brand = document.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = brand.add_run("QuickOps 快维")
    set_run_font(run, 18, True)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("用户使用手册")
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("AI 会话 · 共享终端 · 权限审批 · 主机资产")
    document.add_paragraph()
    scope = document.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(scope, "试用版  |  适用于运维、实施及单机维护人员")
    for _ in range(7):
        document.add_paragraph()
    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(notice, "使用前请核对目标主机、当前权限和待执行操作")
    document.add_page_break()


def add_contents(document: Document) -> None:
    document.add_heading("内容导航", level=1)
    entries = [
        "首次使用与工作区", "AI 会话与文件", "手动命令与共享终端", "四级权限与审批",
        "主机资产与异常守护", "长会话、报告与常见问题", "安全建议与试用反馈",
    ]
    for index, entry in enumerate(entries, 1):
        paragraph = document.add_paragraph(style="List Number")
        add_inline(paragraph, entry)
    tip = document.add_paragraph()
    tip.style = document.styles["Intense Quote"]
    add_inline(tip, "快速开始：登录后先核对右侧目标主机，再到设置中确认模型和工具箱，日常权限建议使用“审批执行”。")


def add_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 4  # Skip title, version and audience lines already represented by the cover.
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("> "):
            paragraph = document.add_paragraph(style="Intense Quote")
            add_inline(paragraph, line[2:])
        elif line.startswith("| "):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [
                [cell.strip() for cell in row.strip("|").split("|")]
                for row in table_lines
                if not re.fullmatch(r"\|?[\s|:-]+\|?", row)
            ]
            if rows:
                table = document.add_table(rows=len(rows), cols=len(rows[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = "Table Grid"
                for row_index, values in enumerate(rows):
                    for column_index, value in enumerate(values):
                        cell = table.cell(row_index, column_index)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        paragraph = cell.paragraphs[0]
                        add_inline(paragraph, value)
                        if row_index == 0:
                            set_cell_shading(cell, TEAL)
                            for run in paragraph.runs:
                                run.font.color.rgb = RGBColor(255, 255, 255)
                                run.bold = True
                        elif row_index % 2 == 0:
                            set_cell_shading(cell, PALE)
                set_repeat_table_header(table.rows[0])
                table.rows[0]._tr.get_or_add_trPr()
            index -= 1
        elif re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
        index += 1


def main() -> None:
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header.paragraphs[0]
    header.text = "QuickOps 快维  |  用户使用手册"
    for run in header.runs:
        set_run_font(run, 8.5, True)
        run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(section.footer.paragraphs[0])

    add_cover(document)
    add_contents(document)
    add_markdown(document, SOURCE.read_text(encoding="utf-8"))

    final = document.add_paragraph()
    final.paragraph_format.space_before = Pt(18)
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = final.add_run("QuickOps 快维 · 让每次操作有证据、可控制、可追溯")
    set_run_font(run, 9.5, True)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    document.core_properties.title = "QuickOps 快维用户使用手册"
    document.core_properties.subject = "QuickOps 试用版操作指南"
    document.core_properties.author = "QuickOps"
    document.core_properties.keywords = "QuickOps, 快维, AI运维, 用户手册"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
