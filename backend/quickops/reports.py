from __future__ import annotations

import html
import json
import re
import shutil
import uuid
from pathlib import Path
from typing import Any

from agno.tools import Toolkit
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.pdfmetrics import registerFont
from reportlab.pdfgen.canvas import Canvas


class ReportArtifactError(ValueError):
    pass


class ReportArtifactStore:
    """Session-isolated downloadable artifacts produced by 小维."""

    formats = {"md", "txt", "html", "json", "docx", "pdf"}
    mime_types = {
        "md": "text/markdown; charset=utf-8",
        "txt": "text/plain; charset=utf-8",
        "html": "text/html; charset=utf-8",
        "json": "application/json",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "pdf": "application/pdf",
    }
    _safe_session = re.compile(r"^[A-Za-z0-9._-]{1,200}$")
    _safe_artifact = re.compile(r"^[0-9a-f]{32}$")

    def __init__(self, root: Path):
        self.root = root.resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def _session_dir(self, session_id: str) -> Path:
        if not self._safe_session.fullmatch(session_id):
            raise ReportArtifactError("无效的会话标识")
        directory = (self.root / session_id).resolve()
        if directory.parent != self.root:
            raise ReportArtifactError("无效的报告目录")
        return directory

    @staticmethod
    def _normalise_format(file_format: str) -> str:
        value = file_format.strip().casefold().lstrip(".")
        aliases = {"markdown": "md", "text": "txt", "word": "docx"}
        return aliases.get(value, value)

    def create(
        self, session_id: str, *, title: str, content: str, file_format: str
    ) -> dict[str, Any]:
        selected = self._normalise_format(file_format)
        if selected not in self.formats:
            raise ReportArtifactError("不支持的报告格式；可选 md、txt、html、json、docx、pdf")
        title = title.strip()[:200] or "QuickOps 运维报告"
        if not content.strip():
            raise ReportArtifactError("报告内容不能为空")
        artifact_id = uuid.uuid4().hex
        directory = self._session_dir(session_id)
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{artifact_id}.{selected}"
        self._write(path, title=title, content=content, file_format=selected)
        metadata = {
            "id": artifact_id,
            "session_id": session_id,
            "title": title,
            "format": selected,
            "mime_type": self.mime_types[selected],
            "size": path.stat().st_size,
            "path": str(path),
        }
        (directory / f"{artifact_id}.json").write_text(
            json.dumps(metadata, ensure_ascii=False), encoding="utf-8"
        )
        return metadata

    def resolve(self, session_id: str, artifact_id: str) -> dict[str, Any]:
        if not self._safe_artifact.fullmatch(artifact_id):
            raise ReportArtifactError("无效的报告标识")
        metadata_path = self._session_dir(session_id) / f"{artifact_id}.json"
        if not metadata_path.is_file():
            raise ReportArtifactError("报告不存在或不属于当前会话")
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        path = Path(metadata["path"]).resolve()
        if path.parent != self._session_dir(session_id) or not path.is_file():
            raise ReportArtifactError("报告文件不可用")
        metadata["path"] = str(path)
        return metadata

    def delete_session(self, session_id: str) -> None:
        directory = self._session_dir(session_id)
        if directory.is_dir():
            shutil.rmtree(directory)

    def _write(self, path: Path, *, title: str, content: str, file_format: str) -> None:
        if file_format == "md":
            body = content if content.lstrip().startswith("#") else f"# {title}\n\n{content}"
            path.write_text(body, encoding="utf-8")
        elif file_format == "txt":
            path.write_text(f"{title}\n\n{self._plain_text(content)}", encoding="utf-8")
        elif file_format == "html":
            escaped = html.escape(content)
            path.write_text(
                '<!doctype html><html lang="zh-CN"><meta charset="utf-8">'
                f"<title>{html.escape(title)}</title><style>body{{font:16px/1.7 system-ui;"
                "max-width:960px;margin:48px auto;padding:0 24px;white-space:pre-wrap;"
                "color:#17242b}}h1{color:#0d7771}</style>"
                f"<body><h1>{html.escape(title)}</h1>{escaped}</body></html>",
                encoding="utf-8",
            )
        elif file_format == "json":
            path.write_text(
                json.dumps({"title": title, "content": content}, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        elif file_format == "docx":
            document = Document()
            document.add_heading(title, level=0)
            for line in content.splitlines():
                stripped = line.strip()
                if not stripped:
                    document.add_paragraph()
                elif stripped.startswith("### "):
                    document.add_heading(stripped[4:], level=3)
                elif stripped.startswith("## "):
                    document.add_heading(stripped[3:], level=2)
                elif stripped.startswith("# "):
                    document.add_heading(stripped[2:], level=1)
                elif stripped.startswith(("- ", "* ")):
                    document.add_paragraph(stripped[2:], style="List Bullet")
                else:
                    document.add_paragraph(stripped)
            document.save(path)
        else:
            self._write_pdf(path, title, self._plain_text(content))

    @staticmethod
    def _plain_text(content: str) -> str:
        content = re.sub(r"```[^\n]*\n?", "", content)
        content = content.replace("```", "")
        return re.sub(r"(?m)^#{1,6}\s+", "", content)

    @staticmethod
    def _write_pdf(path: Path, title: str, content: str) -> None:
        font_name = "STSong-Light"
        registerFont(UnicodeCIDFont(font_name))
        canvas = Canvas(str(path), pagesize=A4)
        width, height = A4
        margin = 48
        y = height - margin
        canvas.setFont(font_name, 18)
        canvas.drawString(margin, y, title)
        y -= 32
        canvas.setFont(font_name, 10.5)
        max_chars = 72
        for paragraph in content.splitlines() or [""]:
            lines = [paragraph[i : i + max_chars] for i in range(0, len(paragraph), max_chars)] or [
                ""
            ]
            for line in lines:
                if y < margin:
                    canvas.showPage()
                    canvas.setFont(font_name, 10.5)
                    y = height - margin
                canvas.drawString(margin, y, line)
                y -= 16
            y -= 4
        canvas.save()


class ReportToolkit(Toolkit):
    def __init__(self, store: ReportArtifactStore, session_id: str) -> None:
        self.store = store
        self.session_id = session_id
        super().__init__(
            name="quickops_reports",
            tools=[self.create_report],
            instructions=(
                "When the operator asks for a report file or specifies a downloadable format, "
                "compose the complete report content and call create_report. Return the supplied "
                "Markdown download link to the operator. Do not substitute chat-only Markdown "
                "for a requested file."
            ),
            add_instructions=True,
        )

    def create_report(self, title: str, content: str, file_format: str = "md") -> str:
        """Create a session-scoped downloadable report in md/txt/html/json/docx/pdf format."""
        metadata = self.store.create(
            self.session_id, title=title, content=content, file_format=file_format
        )
        url = f"/api/quickops/sessions/{self.session_id}/reports/{metadata['id']}"
        return json.dumps(
            {
                "report_id": metadata["id"],
                "title": metadata["title"],
                "format": metadata["format"],
                "size": metadata["size"],
                "download_url": url,
                "download_markdown": f"[下载 {metadata['title']}（.{metadata['format']}）]({url})",
            },
            ensure_ascii=False,
        )
